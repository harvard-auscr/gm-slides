"""Spreadsheet reader: column detection, ordering, cleaning, both file types."""
from __future__ import annotations

import sys
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from shoutout_gen.sheet import (  # noqa: E402
    SheetFormatError, Shoutout, filter_since, is_placeholder, parse_since, parse_timestamp, read_shoutouts,
    split_placeholders,
)


def _csv(tmp_path: Path, header: str, *rows: str) -> Path:
    p = tmp_path / "sheet.csv"
    # newline="" so an embedded "\r\n" inside a quoted cell is written verbatim
    # (Windows text mode would otherwise turn it into "\r\r\n").
    with p.open("w", encoding="utf-8-sig", newline="") as fh:
        fh.write(header + "\n" + "\n".join(rows) + "\n")
    return p


def test_google_form_csv_sorted_by_timestamp(tmp_path):
    p = _csv(
        tmp_path,
        "Timestamp,Shout-out!,Email Address",
        "9/17/2025 20:14:33,second,a@x",
        "9/17/2025 19:01:00,first,b@x",
        "9/17/2025 20:14:33,also second (stable),c@x",
        "9/17/2025 21:00:00,,d@x",
        '9/17/2025 22:00:00,"multi\r\nline",e@x',
    )
    out = read_shoutouts(p)
    assert [s.text for s in out] == ["first", "second", "also second (stable)", "multi\nline"]
    assert out[0].timestamp == datetime(2025, 9, 17, 19, 1)
    assert out[0].row == 3


def test_forms_download_with_email_column_and_question_header(tmp_path):
    """The exact shape of a Forms responses download accepted in wet test C1:
    Timestamp, the collected email, then the question text as the header --
    the email column must never be mistaken for the message."""
    p = _csv(
        tmp_path,
        'Timestamp,Email Address,"Shout-out! Who deserves some love this week, and why?"',
        "9/3/2026 0:35:17,member08@example.edu,thanks for the snacks",
        "9/3/2026 9:06:42,member02@example.edu,the tailgate playlist!!",
    )
    out = read_shoutouts(p)
    assert [s.text for s in out] == ["thanks for the snacks", "the tailgate playlist!!"]
    assert out[0].timestamp == datetime(2026, 9, 3, 0, 35, 17)


def test_untimestamped_rows_go_last_in_sheet_order(tmp_path):
    p = _csv(tmp_path, "Timestamp,Message", "not a date,z", "2025-01-02 10:00,b", ",y", "2025-01-01 10:00,a")
    assert [s.text for s in read_shoutouts(p)] == ["a", "b", "z", "y"]


def test_columns_detected_by_content_when_headers_are_unhelpful(tmp_path):
    p = _csv(tmp_path, "A,B,C", "2025-01-01 10:00,x,a much longer shout-out here", "2025-01-01 11:00,y,another long one")
    out = read_shoutouts(p)
    assert [s.text for s in out] == ["a much longer shout-out here", "another long one"]
    assert out[0].timestamp is not None


def test_forced_columns_and_bad_names(tmp_path):
    p = _csv(tmp_path, "Timestamp,Note,Shoutout", "2025-01-01,keep me,not me")
    assert [s.text for s in read_shoutouts(p, text_column="Note")] == ["keep me"]
    with pytest.raises(SheetFormatError):
        read_shoutouts(p, text_column="Nope")
    with pytest.raises(SheetFormatError):
        read_shoutouts(p, time_column="Nope")


def test_xlsx_with_real_datetime_cells(tmp_path):
    p = tmp_path / "sheet.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Timestamp", "Shout-out"])
    ws.append([datetime(2025, 9, 17, 20, 0), "later"])
    ws.append([datetime(2025, 9, 17, 19, 0), "earlier"])
    ws.append([None, None])
    wb.save(p)
    assert [s.text for s in read_shoutouts(p)] == ["earlier", "later"]


def test_unsupported_and_empty_files(tmp_path):
    with pytest.raises(SheetFormatError):
        read_shoutouts(tmp_path / "x.txt")
    empty = tmp_path / "e.csv"
    empty.write_text("", encoding="utf-8")
    with pytest.raises(SheetFormatError):
        read_shoutouts(empty)


@pytest.mark.parametrize(
    "value, expected",
    [
        ("9/17/2025 20:14:33", datetime(2025, 9, 17, 20, 14, 33)),
        ("2025-09-17T20:14:33", datetime(2025, 9, 17, 20, 14, 33)),
        ("2025-09-17", datetime(2025, 9, 17)),
        (datetime(2025, 1, 1), datetime(2025, 1, 1)),
        ("", None),
        (None, None),
        ("yesterday", None),
    ],
)
def test_parse_timestamp(value, expected):
    assert parse_timestamp(value) == expected


# --------------------------------------------------------------------------- #
# The 2026-27 form and the --since cutoff
# --------------------------------------------------------------------------- #
def test_forms_2026_responses_with_name_and_qotw_columns(tmp_path):
    """The live 2026-27 responses tab (header pulled 2026-09-23): an optional name
    column before the shout-out and a question-of-the-week after it -- neither may
    be mistaken for the message, however wordy the QOTW answers get."""
    p = _csv(
        tmp_path,
        "Timestamp,(Optional) Name,SHOUT \U0001f5e3\ufe0f OUT \U0001f525\U0001f525\U0001f525,"
        "QOTW [9/23]: Which board member is most nonchalant?",
        "9/23/2026 18:02:10,,the tailgate playlist!!,a much longer answer to the question of the week than any shout-out",
        "9/23/2026 16:35:52,member08,thanks for the snacks,someone",
    )
    out = read_shoutouts(p)
    assert [s.text for s in out] == ["thanks for the snacks", "the tailgate playlist!!"]
    assert out[0].timestamp == datetime(2026, 9, 23, 16, 35, 52)


@pytest.mark.parametrize(
    "text, expected",
    [
        ("9.17", datetime(2026, 9, 17)),
        ("9.17 21:00", datetime(2026, 9, 17, 21, 0)),
        ("9.23", datetime(2026, 9, 23)),  # today itself counts
        ("3.30", datetime(2026, 3, 30)),  # 177 days back: still a bare date
        ("2026-09-17", datetime(2026, 9, 17)),
        ("2025-10-01", datetime(2025, 10, 1)),  # a full date reaches back as far as you like
        ("2026-09-17 18:30", datetime(2026, 9, 17, 18, 30)),
        ("9/17/2026 18:30:00", datetime(2026, 9, 17, 18, 30)),
        ("  9.17  ", datetime(2026, 9, 17)),
    ],
)
def test_parse_since(text, expected):
    assert parse_since(text, today=date(2026, 9, 23)) == expected


@pytest.mark.parametrize("bad", ["", "yesterday", "13.40", "9.17.2026", "2.30", "9.17 25:00"])
def test_parse_since_rejects_non_dates(bad):
    with pytest.raises(ValueError):
        parse_since(bad, today=date(2026, 9, 23))


def test_bare_date_after_today_is_last_years_across_new_year():
    assert parse_since("12.10", today=date(2027, 1, 5)) == datetime(2026, 12, 10)


@pytest.mark.parametrize("typo", ["9.24", "10.1", "12.10", "3.20"])
def test_bare_date_more_than_half_a_year_back_is_refused(typo):
    # "9.24" typed on 9/23 would mean 9/24 LAST year and put the whole year on the slide
    with pytest.raises(ValueError, match="give the full date"):
        parse_since(typo, today=date(2026, 9, 23))


@pytest.mark.parametrize("zoned", ["2026-09-17T18:30Z", "2026-09-17 18:30+05:00"])
def test_time_zone_is_refused_not_a_crash_later(zoned):
    with pytest.raises(ValueError, match="time zone"):
        parse_since(zoned, today=date(2026, 9, 23))


def test_filter_since_keeps_the_cutoff_and_untimestamped_rows():
    rows = [
        Shoutout("old", datetime(2026, 9, 10, 19), 2),
        Shoutout("on the cutoff", datetime(2026, 9, 17), 3),
        Shoutout("new", datetime(2026, 9, 20), 4),
        Shoutout("undated", None, 5),
    ]
    kept, skipped = filter_since(rows, datetime(2026, 9, 17))
    assert [s.text for s in kept] == ["on the cutoff", "new", "undated"]
    assert skipped == 1


@pytest.mark.parametrize("text, expected", [
    ("/", True), (".", True), ("-", True), ("??", True), ("/ /", True), ("...", True),
    ("bao", False), ("a/b", False), ("🔥🔥", False), ("??? [tbd]", False), ("7", False),
])
def test_is_placeholder(text, expected):
    """Seen live 2026-09-23: two members typed "/" into the required shout-out box
    to answer the question of the week. Emoji-only messages are real shout-outs."""
    assert is_placeholder(text) is expected


def test_split_placeholders_keeps_order_and_reports_the_rest():
    rows = [Shoutout("/", None, 2), Shoutout("bao", None, 3), Shoutout(".", None, 4), Shoutout("mike", None, 5)]
    kept, dropped = split_placeholders(rows)
    assert [s.text for s in kept] == ["bao", "mike"]
    assert [s.row for s in dropped] == [2, 4]
