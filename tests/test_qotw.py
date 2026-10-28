"""Question-of-the-week tally, chart and CLI (offline)."""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import qotw_chart  # noqa: E402
from shoutout_gen.qotw import pick_column, question_text, read_tally, render_chart, tally  # noqa: E402
from shoutout_gen.sheet import SheetFormatError  # noqa: E402

HEADER = "Timestamp,(Optional) Name,SHOUT OUT,QOTW [9/30]: Which board member is the greenest flag?"


def _csv(tmp_path: Path, *rows: str, header: str = HEADER) -> Path:
    p = tmp_path / "form.csv"
    with p.open("w", encoding="utf-8-sig", newline="") as fh:
        fh.write(header + "\n" + "\n".join(rows) + "\n")
    return p


def test_column_found_by_header_then_letter_then_fallback_d():
    header = HEADER.split(",")
    assert pick_column(header) == header[3]
    assert pick_column(header, "C") == "SHOUT OUT"
    assert pick_column(header, "SHOUT OUT") == "SHOUT OUT"
    assert pick_column(["Timestamp", "a", "b", "Pick one"]) == "Pick one"
    with pytest.raises(SheetFormatError):
        pick_column(header, "Z")


def test_header_that_only_mentions_the_qotw_is_not_it():
    shout = "SHOUT OUT (then answer the QOTW below!)"
    assert pick_column(["Timestamp", shout, "QOTW [10/21]: Who?"]) == "QOTW [10/21]: Who?"
    assert pick_column(["Timestamp", "Name", shout, "Pick one", "Past QOTW answers"]) == "Pick one"  # -> column D


def test_question_text_drops_the_tag():
    assert question_text("QOTW [9/30]: Which board member is the greenest flag?") == "Which board member is the greenest flag?"
    assert question_text("Question of the week - Best snack?") == "Best snack?"
    assert question_text("Who wins?") == "Who wins?"


def test_tally_orders_by_count_then_name_and_merges_spellings():
    answers = [("Bao", None), ("avery l", None), ("Avery L", None), ("Avery L", None), ("Amy", None), ("Bao", None)]
    result = tally("Q?", answers)
    assert result.counts == [("Avery L", 3), ("Bao", 2), ("Amy", 1)]
    assert result.responses == 6


def test_checkbox_cells_count_each_pick_once():
    result = tally("Q?", [("Amy, Bao", None), ("Bao, Bao", None), ("/", None), (", ", None)])
    assert result.counts == [("Bao", 2), ("Amy", 1)]
    assert result.responses == 2  # "/" and ", " picked nobody, so they are not responses


def test_since_keeps_only_this_weeks_answers(tmp_path):
    p = _csv(
        tmp_path,
        "9/23/2026 19:40:00,,hi,Old Pick",
        "9/30/2026 19:00:00,,hi,",
        "10/1/2026 11:57:10,,,Jordan Pike",
        "10/7/2026 10:43:43,,x,Avery Lane",
        "10/7/2026 10:45:59,,y,Avery Lane",
    )
    result, skipped = read_tally(p, since=datetime(2026, 9, 30))
    assert skipped == 1
    assert result.question == "Which board member is the greenest flag?"
    assert result.counts == [("Avery Lane", 2), ("Jordan Pike", 1)]
    assert (result.first, result.last) == (datetime(2026, 10, 1, 11, 57, 10), datetime(2026, 10, 7, 10, 45, 59))


def test_chart_is_a_16_9_png(tmp_path):
    out = render_chart(tally("Q?", [("Amy", None), ("Bao", None)]), tmp_path / "c.png")
    with Image.open(out) as img:
        assert abs(img.width / img.height - 16 / 9) < 0.01


@pytest.mark.parametrize("question", [
    "Which board member is most likely to skip lecture for the DHall on a Monday morning?",
    "If the whole board were stranded on an island for a week, which board member would end up "
    "running the island by Wednesday, and who would they put in charge of the snacks?",
])
def test_long_question_stays_on_the_chart(tmp_path, question):
    out = render_chart(tally(question, [("Amy", None), ("Bao", None)]), tmp_path / "c.png")
    with Image.open(out) as img:
        grey = img.convert("L")
        w, h = grey.size
        right_margin = grey.crop((int(w * 0.965), 0, w, int(h * 0.3)))  # where a clipped title runs off
        assert right_margin.getextrema()[0] >= 250, "title text reaches the right edge"


def test_cli_writes_png_and_one_picture_slide(tmp_path, capsys):
    p = _csv(tmp_path, "10/1/2026 11:57:10,,,Jordan Pike", "10/7/2026 10:43:43,,,Avery Lane")
    code = qotw_chart._main([
        str(p), "--gm", "4", "--date", "10.7", "--since", "2026-09-30",
        "--out", str(tmp_path / "out"), "--env-file", str(tmp_path / "no.env"),
    ])
    assert code == 0
    assert (tmp_path / "out" / "GM #4 10.7---qotw.png").is_file()
    prs = Presentation(str(tmp_path / "out" / "GM #4 10.7---qotw.pptx"))
    assert len(prs.slides) == 1
    [shape] = prs.slides[0].shapes
    assert shape.shape_type == 13  # PICTURE
    assert (shape.width, shape.height) == (prs.slide_width, prs.slide_height)
    assert "2 responses" in capsys.readouterr().out


def test_cli_empty_window_is_an_error(tmp_path, capsys):
    p = _csv(tmp_path, "9/23/2026 19:40:00,,,Old Pick")
    code = qotw_chart._main([
        str(p), "--gm", "4", "--date", "10.7", "--since", "2026-09-30",
        "--out", str(tmp_path / "out"), "--env-file", str(tmp_path / "no.env"),
    ])
    assert code == 2
    assert "no QOTW answers" in capsys.readouterr().err


# --------------------------------------------------------------------------- #
# make_slide.py: shout-outs, then the QOTW pane, then the chart, on one slide
# --------------------------------------------------------------------------- #
def _make_slide(tmp_path, sheet, *extra):
    import make_slide

    return make_slide._main([
        str(sheet), "--gm", "4", "--date", "10.7", "--out", str(tmp_path / "out"),
        "--env-file", str(tmp_path / "no.env"), *extra,
    ])


def _click_targets(slide) -> list[int]:
    """Shape ids in click order, read back from the slide's timing XML."""
    from pptx.oxml.ns import qn

    return [int(el.get("spid")) for el in slide._element.iter(qn("p:spTgt"))]


def test_make_slide_reveals_qotw_after_the_last_shoutout(tmp_path, capsys):
    from shoutout_gen.deck import QOTW_CHART_NAME, QOTW_PANE_NAME

    p = _csv(
        tmp_path,
        "9/30/2026 19:00:00,,old news,",
        "10/1/2026 11:57:10,,first shout,Jordan Pike",
        "10/2/2026 12:00:00,,second shout,",
        "10/7/2026 10:43:43,,third shout,Avery Lane",
    )
    assert _make_slide(tmp_path, p, "--since", "2026-09-30 20:00") == 0
    prs = Presentation(str(tmp_path / "out" / "GM #4 10.7---shoutouts-qotw.pptx"))
    assert len(prs.slides) == 1
    slide = prs.slides[0]
    by_id = {s.shape_id: s for s in slide.shapes}
    order = [by_id[i] for i in _click_targets(slide)]
    assert [s.text_frame.text for s in order[:3]] == ["first shout", "second shout", "third shout"]
    assert [s.name for s in order[3:]] == [QOTW_PANE_NAME, QOTW_CHART_NAME]
    pane, chart = order[3:]
    assert pane.text_frame.text == "QOTW: Which board member is the greenest flag?"
    for shape in (pane, chart):  # both cover the whole slide, so the shout-outs are hidden behind them
        assert (shape.left, shape.top, shape.width, shape.height) == (0, 0, prs.slide_width, prs.slide_height)
    assert (tmp_path / "out" / "GM #4 10.7---qotw.png").is_file()
    assert "(5 clicks)" in capsys.readouterr().out


def test_make_slide_without_answers_is_shoutouts_only(tmp_path, capsys):
    p = _csv(tmp_path, "10/1/2026 11:57:10,,first shout,", "10/2/2026 12:00:00,,second shout,")
    assert _make_slide(tmp_path, p) == 0
    slide = Presentation(str(tmp_path / "out" / "GM #4 10.7---shoutouts-qotw.pptx")).slides[0]
    assert len(_click_targets(slide)) == 2
    assert "no QOTW answers" in capsys.readouterr().out


def test_make_slide_no_qotw_flag(tmp_path):
    p = _csv(tmp_path, "10/1/2026 11:57:10,,first shout,Jordan Pike")
    assert _make_slide(tmp_path, p, "--no-qotw") == 0
    slide = Presentation(str(tmp_path / "out" / "GM #4 10.7---shoutouts-qotw.pptx")).slides[0]
    assert len(_click_targets(slide)) == 1
    assert not (tmp_path / "out" / "GM #4 10.7---qotw.png").exists()
