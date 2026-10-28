"""Snipe slide: reading snipebot's day-by-day workbook, the week and season counts, the charts and the slide (offline)."""
from __future__ import annotations

import os
import sys
import time as clock
from collections import Counter
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import snipe_slide  # noqa: E402
from shoutout_gen.snipes import (  # noqa: E402
    SnipeExportError, end_of_day, find_workbook, fit_text, label_positions, leader_text, load_snipes, rank,
    render_season, render_week, season, text_width_in, week, week_ticks,
)
from shoutout_gen.deck import write_reveal_slide  # noqa: E402

HEADER = ["Date", "Time", "Sniper", "Target", "Sniper key", "Target key", "Sniper points", "Target points"]
# (date, time, sniper, target, sniper points, target points); a selfie's first row carries the photo's extra point.
ROWS = [
    ("2026-09-05", "11:00:00", "Avery Lane", "Jordan Pike", 1, 0),
    ("2026-09-10", "09:00:00", "Jordan Pike", "Avery Lane", 1, 0),
    ("2026-09-16", "23:59:59", "Avery Lane", "Casey Moore", 1, 0),   # last second of a week
    ("2026-09-17", "00:00:00", "Casey Moore", "Avery Lane", 1, 0),   # first second of the next
    ("2026-09-20", "12:00:00", "Riley Shore", "Avery Lane", 2, 1),   # selfie with a sib
    ("2026-09-20", "12:00:00", "Riley Shore", "Jordan Pike", 1, 1),
    ("2026-09-23", "20:00:00", "'=Sam Quill", "Riley Shore", 1, 0),  # formula-guarded name
]


def _workbook(path: Path, rows=ROWS, header=HEADER, sheet="Snipes") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(header)
    keys: dict[str, int] = {}
    for d, t, sniper, target, sp, tp in rows:
        values = {"Date": datetime.fromisoformat(d), "Time": t, "Sniper": sniper, "Target": target,
                  "Sniper key": keys.setdefault(sniper, len(keys) + 1),
                  "Target key": keys.setdefault(target, len(keys) + 1),
                  "Sniper points": sp, "Target points": tp}
        ws.append([values.get(h) for h in header])
    wb.create_sheet("Standings")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


@pytest.fixture
def snipes(tmp_path):
    return load_snipes(_workbook(tmp_path / "fall-2026_day_by_day.xlsx"))


# --------------------------------------------------------------------------- #
# Reading
# --------------------------------------------------------------------------- #
def test_rows_read_by_header_name_in_any_column_order(tmp_path):
    shuffled = HEADER[::-1]
    got = load_snipes(_workbook(tmp_path / "w.xlsx", header=shuffled))
    assert [(s.at, s.sniper, s.target, s.sniper_points, s.target_points) for s in got[:1]] == [
        (datetime(2026, 9, 5, 11), "Avery Lane", "Jordan Pike", 1, 0)]
    assert len(got) == len(ROWS)


def test_formula_guard_is_stripped_from_names(snipes):
    assert snipes[-1].sniper == "=Sam Quill"


def test_blank_rows_are_skipped(tmp_path):
    path = _workbook(tmp_path / "w.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(path)
    wb["Snipes"].append([None] * len(HEADER))
    wb.save(path)
    assert len(load_snipes(path)) == len(ROWS)


def test_not_a_day_by_day_workbook(tmp_path):
    with pytest.raises(SnipeExportError, match="no 'Snipes' sheet"):
        load_snipes(_workbook(tmp_path / "a.xlsx", sheet="Sheet1"))
    with pytest.raises(SnipeExportError, match="Sniper points"):
        load_snipes(_workbook(tmp_path / "b.xlsx", header=[h for h in HEADER if h != "Sniper points"]))


def test_folder_resolves_to_the_newest_workbook_and_skips_office_locks(tmp_path):
    old = _workbook(tmp_path / "spring-2026_day_by_day.xlsx")
    new = _workbook(tmp_path / "fall-2026_day_by_day.xlsx")
    os.utime(old, (clock.time() - 3600, clock.time() - 3600))
    (tmp_path / "~$fall-2026_day_by_day.xlsx").write_bytes(b"lock")
    assert find_workbook(tmp_path) == new
    assert find_workbook(old) == old
    with pytest.raises(FileNotFoundError, match="no such file or folder"):
        find_workbook(tmp_path / "missing")
    (tmp_path / "empty").mkdir()
    with pytest.raises(FileNotFoundError, match="snipebot export"):
        find_workbook(tmp_path / "empty")


# --------------------------------------------------------------------------- #
# Counting
# --------------------------------------------------------------------------- #
def test_rank_ties_share_a_rank_and_the_cut_keeps_a_small_tie():
    board = rank(Counter({"b": 3, "a": 3, "c": 1, "d": 1, "e": 0}), top=3)
    assert [(e.rank, e.name, e.value) for e in board.entries] == [(1, "a", 3), (1, "b", 3), (3, "c", 1), (3, "d", 1)]
    assert (board.hidden, board.hidden_best) == (0, 0)  # zero values never rank
    wide = rank(Counter({f"p{i:02}": 1 for i in range(20)}), top=3)
    assert len(wide.entries) == 3 and (wide.hidden, wide.hidden_best) == (17, 1)


def test_week_is_inclusive_at_both_ends(snipes):
    w = week(snipes, datetime(2026, 9, 10), end_of_day(date(2026, 9, 16)))
    assert w.snipes == 2
    assert [(e.name, e.value) for e in w.made.entries] == [("Avery Lane", 1), ("Jordan Pike", 1)]
    assert w.busiest_day == (date(2026, 9, 10), 1)


def test_week_counts_one_snipe_per_row_regardless_of_points(snipes):
    w = week(snipes, datetime(2026, 9, 17), end_of_day(date(2026, 9, 23)))
    assert (w.snipes, w.snipers, w.people_sniped) == (4, 3, 3)
    assert w.made.entries[0].name == "Riley Shore" and w.made.entries[0].value == 2


def test_empty_week(snipes):
    w = week(snipes, datetime(2026, 10, 1), end_of_day(date(2026, 10, 7)))
    assert (w.snipes, w.made.entries, w.busiest_day) == (0, [], None)


def test_season_points_sum_both_sides_and_stop_at_until(snipes):
    s = season(snipes, end_of_day(date(2026, 9, 20)))
    assert s.snipes == 6 and s.players == 4
    assert {e.name: e.value for e in s.points.entries} == {
        "Avery Lane": 2 + 1, "Riley Shore": 3, "Jordan Pike": 1 + 1, "Casey Moore": 1}
    assert s.sniped.entries[0].name == "Avery Lane" and s.sniped.entries[0].value == 3


def test_race_weeks_count_back_from_until_and_run_cumulatively(snipes):
    s = season(snipes, end_of_day(date(2026, 9, 23)), race_lines=2)
    assert s.week_ends == [date(2026, 9, 9), date(2026, 9, 16), date(2026, 9, 23)]
    assert s.first == date(2026, 9, 5)
    # Avery: 1 by 9/9, +1 on 9/16 23:59:59 (still that week), +1 as Riley's selfie sib on 9/20.
    assert [(line.name, line.values) for line in s.race] == [("Avery Lane", [1, 2, 3]), ("Riley Shore", [0, 0, 3])]


def test_race_stops_mid_day_when_until_has_a_time(snipes):
    s = season(snipes, datetime(2026, 9, 20, 11, 59))
    assert s.week_ends[-1] == date(2026, 9, 20)
    assert {line.name: line.values for line in s.race}["Avery Lane"][-1] == 2  # the 12:00 selfie is after --until


def test_people_sharing_a_display_name_stay_apart_by_key(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Snipes"
    ws.append(HEADER)
    ws.append([datetime(2026, 9, 5), "10:00:00", "Alex Chen", "Jordan Pike", 1, 2, 1, 0])
    ws.append([datetime(2026, 9, 5), "11:00:00", "Alex Chen", "Jordan Pike", 3, 2, 1, 0])
    wb.save(tmp_path / "w.xlsx")
    rows = load_snipes(tmp_path / "w.xlsx")
    w = week(rows, datetime(2026, 9, 1), end_of_day(date(2026, 9, 7)))
    assert [(e.name, e.value) for e in w.made.entries] == [("Alex Chen", 1), ("Alex Chen", 1)]
    assert w.snipers == 2 and season(rows, end_of_day(date(2026, 9, 7))).players == 3


def test_race_keeps_everyone_tied_at_the_cut_up_to_one_line_per_colour(tmp_path):
    rows = [("2026-09-05", f"10:0{i}:00", f"P{i}", "Zed Target", 1, 0) for i in range(6)]
    s = season(load_snipes(_workbook(tmp_path / "w.xlsx", rows=rows)), end_of_day(date(2026, 9, 5)))
    assert len(s.race) == 6 and all(line.values == [1] for line in s.race)
    assert leader_text(s.points, "points") == "6 tied (1 points each)"
    many = [("2026-09-05", f"10:{i:02}:00", f"P{i:02}", "Zed Target", 1, 0) for i in range(12)]
    assert len(season(load_snipes(_workbook(tmp_path / "m.xlsx", rows=many)), end_of_day(date(2026, 9, 5))).race) == 5


def test_leader_text_names_one_or_two_leaders():
    one = rank(Counter({"a": 5, "b": 3}))
    two = rank(Counter({"a": 5, "b": 5, "c": 1}))
    assert leader_text(one, "points") == "a (5 points)"
    assert leader_text(two, "points") == "a and b (5 points each)"
    assert leader_text(rank(Counter()), "points") == ""


def test_unreadable_files_raise_a_clear_error(tmp_path):
    (tmp_path / "snipes.csv").write_text("date,time,sniper,target\n", encoding="utf-8")
    (tmp_path / "~$lock.xlsx").write_bytes(b"\x00lock")
    for bad in ("snipes.csv", "~$lock.xlsx"):
        with pytest.raises(SnipeExportError, match="cannot read"):
            load_snipes(tmp_path / bad)


def test_fractional_points_name_the_row(tmp_path):
    rows = ROWS[:1] + [("2026-09-06", "10:00:00", "Avery Lane", "Jordan Pike", 1.5, 0)]
    with pytest.raises(SnipeExportError, match="row 3"):
        load_snipes(_workbook(tmp_path / "w.xlsx", rows=rows))


def test_week_ticks_thin_out_past_nine_weeks_and_keep_the_newest():
    ends = [date(2026, 9, 9) + timedelta(days=7 * i) for i in range(13)]
    ticks = week_ticks(ends)
    assert len(ticks) <= 9 and ticks[-1] == (12, "12/2")
    assert len(week_ticks(ends[:9])) == 9


def test_end_labels_stay_apart_and_above_the_floor():
    placed = label_positions([100, 3, 2, 2, 1], gap=7.5)
    assert placed[0] == 100 and min(placed) >= 0
    assert all(a - b >= 7.5 - 1e-9 for a, b in zip(placed, placed[1:]))
    assert label_positions([50, 20], gap=5) == [50, 20]


def test_long_names_are_cut_to_fit():
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib import pyplot as plt
    fig = plt.figure(figsize=(13.333, 7.5), dpi=150)
    try:
        short = fit_text(fig, "Avery Lane", 14, 3.0)
        cut = fit_text(fig, "Alexandria Constantinou-Whitfield", 14, 1.5)
        assert short == "Avery Lane"
        assert cut.endswith("…") and text_width_in(fig, cut, 14) <= 1.5
    finally:
        plt.close(fig)


def test_season_needs_at_least_one_snipe(snipes):
    with pytest.raises(ValueError, match="no snipes"):
        season(snipes, datetime(2026, 9, 1))


# --------------------------------------------------------------------------- #
# Drawing and the slide
# --------------------------------------------------------------------------- #
def test_charts_are_16_by_9(tmp_path, snipes):
    until = end_of_day(date(2026, 9, 23))
    for png in (render_week(week(snipes, datetime(2026, 9, 17), until), tmp_path / "w.png"),
                render_week(week(snipes, datetime(2026, 10, 1), end_of_day(date(2026, 10, 7))), tmp_path / "e.png"),
                render_season(season(snipes, until), tmp_path / "s.png")):
        with Image.open(png) as im:
            assert abs(im.width / im.height - 16 / 9) < 0.01


def test_slide_shows_the_week_and_reveals_the_season_on_the_first_click(tmp_path):
    for name in ("a.png", "b.png"):
        Image.new("RGB", (160, 90), "white").save(tmp_path / name)
    out = write_reveal_slide(tmp_path / "a.png", tmp_path / "b.png", tmp_path / "s.pptx", names=("week", "season"))
    prs = Presentation(str(out))
    assert len(prs.slides) == 1
    slide = prs.slides[0]
    pics = list(slide.shapes)
    assert [p.name for p in pics] == ["week", "season"]  # the season draws on top
    assert all((p.left, p.top, p.width, p.height) == (0, 0, prs.slide_width, prs.slide_height) for p in pics)
    xml = slide._element.xml
    targets = [int(v) for v in __import__("re").findall(r'<p:spTgt spid="(\d+)"', xml)]
    assert targets == [pics[1].shape_id]
    assert xml.count('nodeType="clickEffect"') == 1 and 'presetClass="entr"' in xml


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def _run(tmp_path, *args):
    return snipe_slide._main([*args, "--out", str(tmp_path / "out"), "--env-file", str(tmp_path / "absent.env")],
                             today=date(2026, 9, 23))


def test_cli_writes_both_charts_and_the_slide(tmp_path, capsys):
    book = _workbook(tmp_path / "exports" / "fall-2026_day_by_day.xlsx")
    assert _run(tmp_path, str(book.parent), "--gm", "3", "--date", "9.23", "--since", "9.17") == 0
    names = sorted(p.name for p in (tmp_path / "out").iterdir())
    assert names == ["GM #3 9.23---snipes-season.png", "GM #3 9.23---snipes-week.png", "GM #3 9.23---snipes.pptx"]
    out = capsys.readouterr().out
    assert "4 snipes by 3 snipers" in out and "in the lead: Avery Lane and Riley Shore (3 points each)" in out


def test_cli_week_defaults_to_the_seven_days_ending_today(tmp_path, capsys):
    book = _workbook(tmp_path / "w.xlsx")
    assert _run(tmp_path, str(book), "--gm", "3", "--date", "9.23") == 0
    assert "Thu Sep 17 – Wed Sep 23" in capsys.readouterr().out


def test_cli_until_bare_day_means_its_end_and_a_time_is_exact():
    today = date(2026, 10, 28)
    assert snipe_slide.parse_moment("2026-09-16", today, end=True) == datetime.combine(date(2026, 9, 16), time.max)
    assert snipe_slide.parse_moment("10.28", today, end=True) == datetime.combine(today, time.max)
    assert snipe_slide.parse_moment("2026-09-16 19:50", today, end=True) == datetime(2026, 9, 16, 19, 50)
    assert snipe_slide.parse_moment("2026-10-28T19", today, end=True) == datetime(2026, 10, 28, 19)
    assert snipe_slide.parse_moment("10.22", today, end=False) == datetime(2026, 10, 22)


def test_cli_reads_year_less_dates_against_the_given_today(tmp_path, capsys):
    rows = [("2027-09-20", "10:00:00", "Avery Lane", "Jordan Pike", 1, 0)]
    book = _workbook(tmp_path / "w.xlsx", rows=rows)
    assert snipe_slide._main([str(book), "--gm", "1", "--date", "9.22", "--since", "9.16", "--until", "9.22",
                              "--out", str(tmp_path / "out"), "--env-file", str(tmp_path / "absent.env")],
                             today=date(2027, 9, 22)) == 0
    assert "1 snipe by 1 sniper" in capsys.readouterr().out


def test_cli_refuses_a_future_day_that_wraps_to_last_year(tmp_path, capsys):
    book = _workbook(tmp_path / "w.xlsx")
    assert _run(tmp_path, str(book), "--gm", "3", "--date", "9.23", "--since", "9.24") == 2
    err = capsys.readouterr().err
    assert "2025-09-24" in err and "full date" in err
    assert _run(tmp_path, str(book), "--gm", "3", "--date", "9.23", "--since", "2025-09-24") == 0


def test_cli_errors_exit_2_without_writing(tmp_path, capsys):
    book = _workbook(tmp_path / "w.xlsx")
    assert _run(tmp_path, str(book), "--gm", "3", "--date", "9.23", "--since", "2026-09-20", "--until", "2026-09-18") == 2
    assert _run(tmp_path, str(tmp_path / "nowhere"), "--gm", "3", "--date", "9.23") == 2
    assert _run(tmp_path, str(book), "--gm", "3", "--date", "9.1", "--until", "2026-09-01") == 2
    assert _run(tmp_path, str(book), "--gm", "3", "--date", "9.23", "--since", "nonsense") == 2
    (tmp_path / "w.csv").write_text("date,time\n", encoding="utf-8")
    assert _run(tmp_path, str(tmp_path / "w.csv"), "--gm", "3", "--date", "9.23") == 2
    assert not (tmp_path / "out").exists()
    err = capsys.readouterr().err
    assert "2026-09-20 00:00" in err and "2026-09-18 23:59" in err and "cannot read w.csv" in err
